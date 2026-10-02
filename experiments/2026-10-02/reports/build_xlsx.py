import pandas as pd, numpy as np, json
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
R4='/home/claude/jss_run4/dash_replication_results.csv'; R3='/home/claude/jss_run3/dash_paper_results.csv'; R5='/home/claude/jss_run5/dash_robustness_results.csv'
def rd(p, sections):
    it = pd.read_csv(p, low_memory=False, dtype={'text_value':str,'success':str,'stage2_fired':str,'label_match':str}, chunksize=200000)
    return pd.concat([ch[ch.section.isin(sections)] for ch in it], ignore_index=True)
c4 = rd(R4, ['01_environment','05_prediction_equivalence','08_latency_run_summary','09_latency_condition_summary','10_deadline_coverage'])
c4['env'] = c4.notes.str.extract(r'env=(\w+)')[0]
S = {}
# README
S['README'] = pd.DataFrame({'Item':[
 'Purpose','Sheets','Provenance','Raw data','Time unit','Statistics','Important'],
 'Description':[
 'All DASH results produced on 2026-10-02 for the JSS/ICPE submission, summarised. Raw per-request rows remain in the CSV files (Excel row limit).',
 'Environments, Replication_Summary, Replication_Paired (formulas), Decisions, E0_Stage2, Robustness_W, Robustness_CPU, Proxy_Matrix, Network_Calibration, Lua_Cost, Gating, E2_Memory, Lua_Correctness, Historical_E6, Limitations',
 'cloud = rep_cloud_20261002T141735Z (2 vCPU cloud container); laptop = rep_laptop_20261002T201141Z (Intel Core 5 210H, WSL2). Robustness = robust_cloud_20261002 (cloud, single host). Proxy_Matrix/Lua_Cost/Gating = dash3_20261002T1122Z (cloud, via delay proxy, default XGBoost threading). E2/Lua correctness = dash_20261002T110902Z.',
 'dash_replication_results.csv, dash_robustness_results.csv, dash_paper_results.csv',
 'ms for latency, bytes for memory',
 'Quantiles: numpy linear (Hyndman-Fan type 7). Five repeats per condition; values are descriptive; no confidence intervals unless stated as exploratory.',
 'Environments are never pooled. Proxy_Matrix used a different network path and thread setting from the P0-direct replication and is not directly comparable.']})
# Environments
env = c4[c4.section=='01_environment']; env = env.assign(value_all=env.text_value.fillna(env.value.astype(str)))
S['Environments'] = env.pivot_table(index='metric', columns='env', values='value_all', aggfunc='first').reset_index()
_lv = {'numpy':'2.5.3','pandas':'3.0.6','sklearn':'1.6.1','xgboost':'3.4.1','redis_py':'8.1.0','torch':'2.14.1+cpu','shared_cloud_host':'no: author laptop, Windows + WSL2 Ubuntu (Python 3.13 venv via uv); other apps closed, mains power','host_cpu_monitor':'same as cloud','python_readme_target':'3.11 (not met; 3.13 used)'}
_e = S['Environments']; _e['laptop'] = [ (_lv[m] + (' (from uv install log)' if m in ('numpy','pandas','sklearn','xgboost','redis_py','torch') else '')) if (pd.isna(v) and m in _lv) else v for m, v in zip(_e.metric, _e.laptop)]
# Replication summary
cs = c4[c4.section=='09_latency_condition_summary']
t = cs[cs.metric.isin(['latency_p50','latency_p99','latency_mean','host_cpu_percent_mean','failure_rate']) & cs.aggregation.isin(['mean_of_repeat_values','sd_of_repeat_values','min_of_repeat_values','max_of_repeat_values'])]
t = t.pivot_table(index=['env','K','engine'], columns=['metric','aggregation'], values='value').round(4)
t.columns = [f'{m} | {a.replace("_of_repeat_values","")}' for m,a in t.columns]; S['Replication_Summary'] = t.reset_index()
# Paired (with formulas later)
rs = c4[c4.section=='08_latency_run_summary']
p = rs[rs.metric.isin(['latency_p99','latency_p50']) & rs.engine.isin(['A_ProposedStateful','B1_RedisFetch'])].pivot_table(index=['env','K','repeat'], columns=['metric','engine'], values='value')
p.columns = [f'{m}_{e.split("_")[0]}' for m,e in p.columns]; p = p.reset_index()[['env','K','repeat','latency_p99_A','latency_p99_B1','latency_p50_A','latency_p50_B1']].round(4)
S['Replication_Paired'] = p
# Decisions
d = cs[cs.experiment=='preregistered_decision'][['env','K','metric','value','text_value','status','notes']]
S['Decisions'] = d
# E0
e = c4[(c4.section=='05_prediction_equivalence') & (c4.record_type=='run_summary')]
e = e.assign(env=e.run_id.str.extract(r'rep_(\w+?)_')[0])
S['E0_Stage2'] = e.pivot_table(index=['env','model','engine'], columns='metric', values='value', aggfunc='first').reset_index()
# Robustness
c5 = rd(R5, ['09_latency_condition_summary'])
w = c5[(c5.experiment=='robustness_W_sweep') & (c5.aggregation.isin(['mean_of_repeat_values','sd_of_repeat_values'])) & c5.metric.isin(['state_ms_mean','latency_p50','latency_p99','p99_ratio_B1_over_A'])]
w = w.pivot_table(index=['W','engine'], columns=['metric','aggregation'], values='value').round(4); w.columns=[f'{m} | {a.replace("_of_repeat_values","")}' for m,a in w.columns]
S['Robustness_W'] = w.reset_index()
cp = c5[(c5.experiment=='robustness_cpu_allocation') & (c5.aggregation.isin(['mean_of_repeat_values','median_of_repeat_values','min_of_repeat_values','max_of_repeat_values'])) & c5.metric.isin(['latency_p50','latency_p99','p99_ratio_B1_over_A'])]
cp = cp.assign(cpu=cp.notes.str[10:]).pivot_table(index=['cpu','engine'], columns=['metric','aggregation'], values='value').round(4); cp.columns=[f'{m} | {a.replace("_of_repeat_values","")}' for m,a in cp.columns]
S['Robustness_CPU'] = cp.reset_index()
# prior run3
c3 = rd(R3, ['04_memory_scaling','03_lua_correctness','06_network_calibration','09_latency_condition_summary','11_lua_update_cost','12_gating_ablation','13_historical_results','10_deadline_coverage'])
pm = c3[(c3.section=='09_latency_condition_summary') & c3.aggregation.isin(['mean_of_repeat_values','sd_of_repeat_values']) & c3.metric.isin(['latency_p50','latency_p99','p99_ratio_B1_over_A'])]
pm = pm.pivot_table(index=['network_profile','K','engine'], columns=['metric','aggregation'], values='value').round(4); pm.columns=[f'{m} | {a.replace("_of_repeat_values","")}' for m,a in pm.columns]
S['Proxy_Matrix'] = pm.reset_index()
nc = c3[(c3.section=='06_network_calibration') & (c3.record_type=='run_summary')]
S['Network_Calibration'] = nc.pivot_table(index=['policy'], columns='metric', values='value').round(4).reset_index().rename(columns={'policy':'calibration_tag'})
lc = c3[(c3.section=='11_lua_update_cost') & (c3.record_type=='condition_summary')]
S['Lua_Cost'] = lc.pivot_table(index='update_method', columns=['metric','aggregation'], values='value').round(5).pipe(lambda x: x.set_axis([f'{m} | {a.replace("_of_repeat_values","")}' for m,a in x.columns], axis=1)).reset_index()
g = c3[(c3.section=='12_gating_ablation') & (c3.record_type=='condition_summary') & (c3.aggregation=='mean_of_repeat_values') & (c3.metric.isin(['stage2_fire_rate','latency_p50','latency_p99','latency_mean','prediction_differs_rate']))]
S['Gating'] = g.pivot_table(index=['network_profile','policy'], columns='metric', values='value').round(4).reset_index()
m = c3[(c3.section=='04_memory_scaling') & (c3.record_type=='condition_summary')]
S['E2_Memory'] = m.pivot_table(index=['W','N','logical_payload_bytes'], columns='aggregation', values='value').reset_index()
S['Lua_Correctness'] = c3[c3.section=='03_lua_correctness'][['provenance','update_method','W','N','metric','text_value','status','notes']]
h = c3[(c3.section=='13_historical_results') & (c3.metric=='deadline_coverage') & (c3.K==1)]
S['Historical_E6'] = h.pivot_table(index=['source_file','model','engine'], columns='deadline_ms', values='value').round(4).reset_index()
S['Limitations'] = pd.DataFrame({'#': range(1,15), 'Limitation': [
 'Two environments only (cloud container + one laptop under WSL2); not a homogeneous multi-instance replication.',
 'Environments differ in Redis version (7.0.15 vs 8.0.5), CPU (2 vCPU Xeon vs 12-thread Core 5 210H) and virtualisation; reported separately, never pooled.',
 'K=1: direction NOT maintained across environments (cloud median P99 ratio 1.28, laptop 0.89); scope rule applies on the laptop: benefit limited without added network delay.',
 'K=100: direction maintained (cloud 1.17, laptop 1.43).',
 'Laptop P99 is noisy (A P99 SD 1.76 ms at K=1); noise gate passed (B2 CV 0.09 and 0.23) but K=100 is close to the 0.25 threshold.',
 'P50 ratios (B1/A > 1 in all 20 repeats across both environments) are descriptive and were not a pre-registered decision metric.',
 'Closed-loop single client; K replicated streams of one symbol; no multi-client, arrival-rate, queueing or capacity claims.',
 'Five repeats per condition; P99 of 5,000 requests rests on ~50 observations; no dependence-aware CIs.',
 'Network delay results (Proxy_Matrix) used a user-space Python proxy (netem unavailable) and default XGBoost threading; synthetic, not comparable to the P0-direct runs.',
 'Robustness (W, CPU) results come from one host; the 1-core condition restricts only the client process.',
 'E0 windows were rule-selected to exercise Stage 2; their firing rates are not representative; agreement is between execution paths, not accuracy.',
 'Historical E6 logs (tau=0.6731, older code) are separate provenance and not comparable.',
 'Bounded state: fixed W,d -> independent of N, logical payload 4*d*min(N,W) bytes; not W-independent.',
 'Python 3.13 in both environments (README target 3.11); scikit-learn 1.6.1 met.']})
out = '/home/claude/final/DASH_all_results_20261002.xlsx'
with pd.ExcelWriter(out, engine='openpyxl') as xw:
    for n, df in S.items(): df.to_excel(xw, sheet_name=n, index=False)
wb = load_workbook(out)
ws = wb['Replication_Paired']; n = ws.max_row
for col, head in ((8, 'p99_ratio_B1_over_A'), (9, 'p99_diff_B1_minus_A'), (10, 'p50_ratio_B1_over_A')): ws.cell(1, col, head)
for r in range(2, n + 1):
    ws.cell(r, 8, f'=E{r}/D{r}'); ws.cell(r, 9, f'=E{r}-D{r}'); ws.cell(r, 10, f'=G{r}/F{r}')
    for col in (8, 9, 10): ws.cell(r, col).number_format = '0.000'
ws.cell(n + 2, 1, 'Ratios/differences are formulas over the per-repeat P99/P50 values (ms) in columns D–G, measured per environment.')
hdr = PatternFill('solid', fgColor='DCE6F1')
for w_ in wb.worksheets:
    for row in w_.iter_rows():
        for cell in row: cell.font = Font(name='Arial', size=10, bold=(cell.row == 1))
    for cell in w_[1]: cell.fill = hdr; cell.alignment = Alignment(wrap_text=True, vertical='top')
    w_.freeze_panes = 'A2'
    for i, col in enumerate(w_.columns, 1):
        L = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        w_.column_dimensions[get_column_letter(i)].width = min(max(10, L + 2), 60 if w_.title in ('README', 'Limitations', 'Decisions', 'Lua_Correctness') else 28)
for t_ in ('README', 'Limitations'):
    for row in wb[t_].iter_rows(min_row=2):
        for cell in row: cell.alignment = Alignment(wrap_text=True, vertical='top')
    wb[t_].column_dimensions['B'].width = 110
wb.save(out); print({k: v.shape for k, v in S.items()})
